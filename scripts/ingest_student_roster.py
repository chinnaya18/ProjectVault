import os
import sys
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import psycopg2

# Configure UTF-8 stdout for Windows console
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EXCEL_PATH = r"C:\Users\Balasubramaniam\Downloads\student_list.xlsx"
BACKUP_EXCEL_PATH = r"C:\Users\Balasubramaniam\Downloads\student_list_updated.xlsx"
REPO_DATA_DIR = r"d:\projectvault\data"
REPO_EXCEL_PATH = os.path.join(REPO_DATA_DIR, "student_list_25mx.xlsx")

DB_PARAMS = {
    "dbname": "projectvault",
    "user": "postgres",
    "password": "0608",
    "host": "localhost",
    "port": 5432
}

def process_roster():
    print(f"Loading Excel file: {EXCEL_PATH}")
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Excel file not found at {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.active

    # Check header
    header_sno = sheet.cell(1, 1).value
    header_roll = sheet.cell(1, 2).value
    header_name = sheet.cell(1, 3).value
    print(f"Original Headers: {[header_sno, header_roll, header_name]}")

    # Set Header for Column 4: Email
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    cell_email_header = sheet.cell(1, 4, value="Email")
    cell_email_header.font = header_font
    cell_email_header.fill = header_fill
    cell_email_header.alignment = center_align

    # Also style first 3 headers cleanly if needed
    for col in range(1, 4):
        sheet.cell(1, col).font = header_font
        sheet.cell(1, col).fill = header_fill
        sheet.cell(1, col).alignment = center_align

    students = []
    max_row = sheet.max_row

    for r in range(2, max_row + 1):
        sno = sheet.cell(r, 1).value
        roll_raw = sheet.cell(r, 2).value
        name_raw = sheet.cell(r, 3).value

        if not roll_raw or not name_raw:
            continue

        roll_no = str(roll_raw).strip().upper()
        name = str(name_raw).strip()
        email = f"{roll_no.lower()}@psgtech.ac.in"

        # Write email into Excel column 4
        cell_email = sheet.cell(r, 4, value=email)
        cell_email.alignment = left_align
        cell_email.font = Font(name="Calibri", size=11)

        students.append({
            "sno": sno,
            "roll_no": roll_no,
            "name": name,
            "email": email
        })

    # Adjust column widths
    sheet.column_dimensions['A'].width = 8
    sheet.column_dimensions['B'].width = 16
    sheet.column_dimensions['C'].width = 32
    sheet.column_dimensions['D'].width = 30

    # Save to original and backups
    print(f"Extracted {len(students)} student records.")
    wb.save(EXCEL_PATH)
    print(f"Updated original Excel: {EXCEL_PATH}")

    wb.save(BACKUP_EXCEL_PATH)
    print(f"Saved backup Excel: {BACKUP_EXCEL_PATH}")

    os.makedirs(REPO_DATA_DIR, exist_ok=True)
    wb.save(REPO_EXCEL_PATH)
    print(f"Saved repository copy: {REPO_EXCEL_PATH}")

    # Now ingest into PostgreSQL student_roster table
    print("\nConnecting to PostgreSQL...")
    conn = psycopg2.connect(**DB_PARAMS)
    cur = conn.cursor()

    # Insert / Upsert into student_roster
    upsert_sql = """
        INSERT INTO student_roster (roll_no, name, email, batch, department_id, is_eligible, is_registered, created_at, updated_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, NOW(), NOW())
        ON CONFLICT (roll_no) DO UPDATE SET
            name = EXCLUDED.name,
            email = EXCLUDED.email,
            batch = EXCLUDED.batch,
            department_id = EXCLUDED.department_id,
            is_eligible = EXCLUDED.is_eligible,
            updated_at = NOW();
    """

    for s in students:
        cur.execute(upsert_sql, (
            s["roll_no"],
            s["name"],
            s["email"],
            "2025-2027",
            1, # MCA Department
            True, # is_eligible
            False # is_registered
        ))

    conn.commit()

    # Query verification stats
    cur.execute("SELECT COUNT(*) FROM student_roster WHERE batch = '2025-2027' AND department_id = 1;")
    total_mca = cur.fetchone()[0]

    cur.execute("SELECT roll_no, name, email FROM student_roster WHERE batch = '2025-2027' ORDER BY roll_no ASC LIMIT 5;")
    sample_first = cur.fetchall()

    cur.execute("SELECT roll_no, name, email FROM student_roster WHERE batch = '2025-2027' ORDER BY roll_no DESC LIMIT 5;")
    sample_last = cur.fetchall()

    cur.close()
    conn.close()

    print("\n--- INGESTION SUMMARY ---")
    print(f"Total students ingested into student_roster: {total_mca}")
    print("\nFirst 5 Students in Roster:")
    for roll, name, email in sample_first:
        print(f"  [{roll}] {name} -> {email}")
    print("\nLast 5 Students in Roster:")
    for roll, name, email in sample_last:
        print(f"  [{roll}] {name} -> {email}")

if __name__ == "__main__":
    process_roster()
